"""
Attendance Manager Module
=========================
Manages course-level attendance sessions with P/A/OD/ML codes,
live calculations, Excel import/export in competition format.

Classes
-------
AttendanceSession  -- one session per course per date

Functions
---------
load_course_history(program, semester, section, course, attendance_dir)
"""

import os, io
from datetime import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from config import BASE_DIR, ATTENDANCE_DIR

ATTENDANCE_CODES = {
    "P":  "Present",
    "A":  "Absent",
    "OD": "On Duty / Event",
    "ML": "Medical Leave",
}

CODE_COLORS = {
    "P":  "C8E6C9",
    "A":  "FFCDD2",
    "OD": "FFF9C4",
    "ML": "E3F2FD",
    "":   "F5F5F5",
}


def _thin_side():
    return Side(style="thin")

def _thin_border():
    s = _thin_side()
    return Border(left=s, right=s, top=s, bottom=s)

def _center():
    return Alignment(horizontal="center", vertical="center")

def _left():
    return Alignment(horizontal="left", vertical="center")

def _fill(hex_color):
    return PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")


class AttendanceSession:
    """Manages attendance for one course session.

    Parameters
    ----------
    program   : str  -- e.g. "BCA"
    semester  : str  -- e.g. "Semester 3"
    section   : str  -- e.g. "A"
    course    : str  -- e.g. "Database Management Systems"
    students  : list -- [{roll_no, enrollment_no, name}, ...]
    date      : str  -- YYYY-MM-DD (defaults to today)
    marks     : dict -- {roll_no: code} to pre-populate marks (optional)
    """

    def __init__(self, program, semester, section, course,
                 students, date=None, marks=None):
        self.program  = program
        self.semester = semester
        self.section  = section
        self.course   = course
        self.date     = date or datetime.now().strftime("%Y-%m-%d")
        self.time     = datetime.now().strftime("%H:%M:%S")
        self.students = [dict(s) for s in students]
        self._marks   = {}
        for s in self.students:
            self._marks[s["roll_no"]] = "A"
        if marks:
            for roll, code in marks.items():
                if roll in self._marks and code in ATTENDANCE_CODES:
                    self._marks[roll] = code

    # -- Marking ---

    def mark(self, roll_no, code):
        """Mark a student. Returns False if roll_no not in session."""
        code = str(code).upper().strip()
        if code not in ATTENDANCE_CODES:
            return False
        if roll_no not in self._marks:
            return False
        self._marks[roll_no] = code
        return True

    set_mark = mark  # Alias for compatibility

    def mark_by_name(self, name, code):
        """Mark by student name (case-insensitive). Returns True if found."""
        for s in self.students:
            if s["name"].lower() == name.lower():
                return self.mark(s["roll_no"], code)
        return False

    def mark_event(self, roll_no, event_name="Approved Event"):
        """Record event/activity attendance for a student: marks OD and records event."""
        if roll_no not in self._marks:
            return False
        self._marks[roll_no] = "OD"
        if not hasattr(self, "_student_events"):
            self._student_events = {}
        self._student_events[roll_no] = event_name
        try:
            import academic_db
            sname = next((s["name"] for s in self.students if s["roll_no"] == roll_no), "")
            academic_db.record_event_attendance(event_name, roll_no, sname)
        except Exception:
            pass
        return True

    def get_mark(self, roll_no):
        """Return current mark for roll_no, or '' if not found."""
        return self._marks.get(roll_no, "")

    def get_all_marks(self):
        """Return {roll_no: code} dict."""
        return dict(self._marks)

    # -- Statistics ---

    def get_stats(self):
        """Return live session statistics dict."""
        marks  = list(self._marks.values())
        total  = len(marks)
        pres   = marks.count("P")
        abst   = marks.count("A")
        od     = marks.count("OD")
        ml     = marks.count("ML")
        pct    = round((pres + od) / total * 100, 1) if total > 0 else 0.0
        return {
            "total": total, "conducted": 1,
            "present": pres, "absent": abst,
            "od": od, "ml": ml, "pct": pct,
        }

    # -- Excel Export ---

    def export_to_excel(self, file_path=None):
        """Export competition-format attendance sheet. Returns saved file path."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Attendance"

        hdr_font  = Font(name="Calibri", bold=True, size=12, color="FFFFFF")
        hdr_fill  = _fill("1a1a2e")
        ca        = _center()
        la        = _left()
        bdr       = _thin_border()

        # Row 1: system header
        ws.merge_cells("A1:J1")
        ws["A1"] = "Smart Attendance Management System"
        ws["A1"].font = Font(name="Calibri", bold=True, size=14, color="FFFFFF")
        ws["A1"].fill = _fill("0f3460")
        ws["A1"].alignment = ca
        ws.row_dimensions[1].height = 30

        # Rows 2-5: meta block
        meta_fill = _fill("16213e")
        meta_font = Font(name="Calibri", size=11, color="e0e0e0")
        bold_font = Font(name="Calibri", bold=True, size=11, color="e0e0e0")
        meta_rows = [
            ("Program:", self.program, "Semester:", self.semester),
            ("Section:", self.section, "Course:", self.course),
            ("Date:", self.date, "Time:", self.time),
            ("Faculty:", "", "Academic Year:", ""),
        ]
        for r, (k1, v1, k2, v2) in enumerate(meta_rows, 2):
            ws.merge_cells("A" + str(r) + ":B" + str(r))
            ws.merge_cells("C" + str(r) + ":E" + str(r))
            ws.merge_cells("F" + str(r) + ":G" + str(r))
            ws.merge_cells("H" + str(r) + ":J" + str(r))
            for col_letter, value, fnt in [
                ("A", k1, bold_font), ("C", v1, meta_font),
                ("F", k2, bold_font), ("H", v2, meta_font)
            ]:
                cell = ws[col_letter + str(r)]
                cell.value = value
                cell.font  = fnt
                cell.fill  = meta_fill
                cell.alignment = la
            ws.row_dimensions[r].height = 20

        # Row 6: column headers
        col_headers = [
            "S.No", "Roll No", "Enrollment No", "Student Name",
            "Status", "Total Conducted", "Classes Attended",
            "Classes Absent", "Activity Engagement", "Attendance %"
        ]
        col_widths = [7, 12, 16, 25, 10, 16, 17, 15, 19, 14]
        for c, (hdr, width) in enumerate(zip(col_headers, col_widths), 1):
            cell = ws.cell(row=6, column=c, value=hdr)
            cell.font      = hdr_font
            cell.fill      = hdr_fill
            cell.alignment = ca
            cell.border    = bdr
            ws.column_dimensions[get_column_letter(c)].width = width
        ws.row_dimensions[6].height = 22

        # Row 7+: student rows
        for row_idx, student in enumerate(self.students, 7):
            roll   = student["roll_no"]
            enrl   = student.get("enrollment_no", "")
            name   = student["name"]
            code   = self._marks.get(roll, "A")
            rfill  = _fill(CODE_COLORS.get(code, "F5F5F5"))
            attended = 1 if code in ("P", "OD") else 0
            absent   = 1 if code == "A"        else 0
            activity = 1 if code == "OD"       else 0
            pct_val  = 100.0 if attended else 0.0
            row_data = [row_idx - 6, roll, enrl, name, code,
                        1, attended, absent, activity, pct_val]
            for col_idx, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.fill      = rfill
                cell.border    = bdr
                cell.alignment = la if col_idx == 4 else ca
                if col_idx == 10:
                    cell.number_format = "0.0\"%\""
            ws.row_dimensions[row_idx].height = 18

        ws.freeze_panes = "A7"

        # Summary row
        sum_row = len(self.students) + 7
        stats = self.get_stats()
        sfill = _fill("1a1a2e")
        sfont = Font(name="Calibri", bold=True, color="FFFFFF")

        cell_summary = ws.cell(row=sum_row, column=1, value="SUMMARY")
        cell_summary.font      = sfont
        cell_summary.fill      = sfill
        cell_summary.alignment = ca
        cell_summary.border    = bdr
        ws.merge_cells("A" + str(sum_row) + ":D" + str(sum_row))

        sum_data = [
            "Total: " + str(stats["total"]),
            stats["conducted"],
            stats["present"],
            stats["absent"],
            stats["od"],
            str(stats["pct"]) + "%"
        ]
        for col_idx, value in enumerate(sum_data, 5):
            cell = ws.cell(row=sum_row, column=col_idx, value=value)
            cell.font      = sfont
            cell.fill      = sfill
            cell.alignment = ca
            cell.border    = bdr

        if not file_path:
            safe_course = self.course.replace(" ", "_").replace("/", "-")
            safe_prog   = self.program.replace(".", "").replace(" ", "_")
            safe_sem    = self.semester.replace(" ", "_")
            fname = "attendance_" + safe_prog + "_" + safe_sem + "_" + self.section + "_" + safe_course + "_" + self.date + ".xlsx"
            file_path = os.path.join(ATTENDANCE_DIR, fname)
        os.makedirs(os.path.dirname(file_path), exist_ok=True)
        wb.save(file_path)
        return file_path

    def export_to_bytes(self):
        """Export to bytes (for Streamlit download button)."""
        import tempfile
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            tmp_path = tmp.name
        try:
            self.export_to_excel(tmp_path)
            with open(tmp_path, "rb") as fh:
                data = fh.read()
        finally:
            try:
                os.unlink(tmp_path)
            except Exception:
                pass
        return data

    # -- Excel Import ---

    def import_from_excel(self, file_path_or_bytes):
        """Import marks from existing Excel. Returns (merged, skipped, message)."""
        try:
            if isinstance(file_path_or_bytes, (str, os.PathLike)):
                wb = load_workbook(file_path_or_bytes, read_only=True, data_only=True)
            elif isinstance(file_path_or_bytes, bytes):
                wb = load_workbook(io.BytesIO(file_path_or_bytes), read_only=True, data_only=True)
            else:
                wb = load_workbook(file_path_or_bytes, read_only=True, data_only=True)

            ws = wb.active
            merged = skipped = 0
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or not row[1]:
                    continue
                roll_val = str(row[1]).strip()
                status_val = ""
                for col_idx in [4, 5, 6, 7]:
                    if len(row) > col_idx and row[col_idx]:
                        val = str(row[col_idx]).strip().upper()
                        if val == "PRESENT": val = "P"
                        if val == "ABSENT":  val = "A"
                        if val in ATTENDANCE_CODES:
                            status_val = val
                            break
                if roll_val and status_val and roll_val in self._marks:
                    self._marks[roll_val] = status_val
                    merged += 1
                else:
                    skipped += 1
            wb.close()
            return merged, skipped, "Imported: " + str(merged) + " marks merged, " + str(skipped) + " rows skipped."
        except Exception as exc:
            return 0, 0, "Import failed: " + str(exc)

    def import_from_bytes(self, data_bytes):
        """Convenience wrapper for import_from_excel with bytes input."""
        return self.import_from_excel(data_bytes)


# -- Module-level helper ---

def load_course_history(program, semester, section, course, attendance_dir=None):
    """Scan attendance_dir for files matching this course. Returns list of {date, file_path}."""
    if attendance_dir is None:
        attendance_dir = ATTENDANCE_DIR
    sessions = []
    safe_course = course.replace(" ", "_").replace("/", "-")
    safe_prog   = program.replace(".", "").replace(" ", "_")
    safe_sem    = semester.replace(" ", "_")
    prefix = "attendance_" + safe_prog + "_" + safe_sem + "_" + section + "_" + safe_course + "_"
    if not os.path.exists(attendance_dir):
        return sessions
    for fname in sorted(os.listdir(attendance_dir)):
        if fname.startswith(prefix) and fname.endswith(".xlsx"):
            date_part = fname[len(prefix):].replace(".xlsx", "")
            sessions.append({"date": date_part, "file_path": os.path.join(attendance_dir, fname)})
    return sessions
