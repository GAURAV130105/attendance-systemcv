import os, json, threading, io
from config import BASE_DIR

ACADEMIC_DATA_FILE = os.path.join(BASE_DIR, "academic_data.json")
_lock = threading.Lock()

_DEFAULT_DATA = {
    "programs": {
        "BCA": {
            "semesters": {
                "Semester 1": {"sections": ["A","B"], "courses": ["Computer Fundamentals","Mathematics I","English Communication","Programming in C","Digital Electronics"]},
                "Semester 2": {"sections": ["A","B"], "courses": ["Data Structures","Mathematics II","OOP with C++","Web Technologies","PC Hardware and Networking"]},
                "Semester 3": {"sections": ["A","B"], "courses": ["Database Management Systems","Operating Systems","JAVA Programming","Computer Networks","Software Engineering"]},
                "Semester 4": {"sections": ["A","B"], "courses": ["Advanced Java","System Analysis and Design","Computer Graphics","Python Programming","DBMS Lab"]},
                "Semester 5": {"sections": ["A"],     "courses": ["Machine Learning","Cloud Computing","Mobile Computing","Cyber Security","Elective I"]},
                "Semester 6": {"sections": ["A"],     "courses": ["Project Work","Entrepreneurship","Artificial Intelligence","Elective II","Internship"]}
            }
        },
        "MCA": {
            "semesters": {
                "Semester 1": {"sections": ["A"], "courses": ["Discrete Mathematics","Computer Organization","Programming with Python","Data Structures and Algorithms","Technical Writing"]},
                "Semester 2": {"sections": ["A"], "courses": ["Design and Analysis of Algorithms","Database Management Systems","Advanced OOP","Computer Networks","Probability and Statistics"]},
                "Semester 3": {"sections": ["A"], "courses": ["Software Engineering","Machine Learning","Operating Systems","Distributed Systems","Web Services and API"]},
                "Semester 4": {"sections": ["A"], "courses": ["Project","Internship","Research Methodology","Cloud Computing","Elective"]}
            }
        },
        "B.Tech CSE": {
            "semesters": {
                "Semester 1": {"sections": ["A","B","C"], "courses": ["Engineering Mathematics I","Engineering Physics","Engineering Chemistry","Basic Electronics","Programming in C"]},
                "Semester 2": {"sections": ["A","B","C"], "courses": ["Engineering Mathematics II","Engineering Mechanics","Electrical Engineering","Data Structures","OOP with Java"]},
                "Semester 3": {"sections": ["A","B","C"], "courses": ["Discrete Mathematics","Digital Logic Design","Database Management","Computer Architecture","OOP with C++"]},
                "Semester 4": {"sections": ["A","B","C"], "courses": ["Operating Systems","Computer Networks","Theory of Computation","Microprocessors","Software Engineering"]},
                "Semester 5": {"sections": ["A","B"],     "courses": ["Compiler Design","Artificial Intelligence","Machine Learning","Web Technologies","Elective I"]},
                "Semester 6": {"sections": ["A","B"],     "courses": ["Cloud Computing","Information Security","Big Data Analytics","Mobile Computing","Elective II"]},
                "Semester 7": {"sections": ["A"],         "courses": ["Project Phase I","Advanced Machine Learning","Internet of Things","Elective III","Industrial Training"]},
                "Semester 8": {"sections": ["A"],         "courses": ["Project Phase II","Entrepreneurship","Research Project","Elective IV"]}
            }
        }
    },
    "students": {}
}


def _load_data():
    with _lock:
        if not os.path.exists(ACADEMIC_DATA_FILE):
            with open(ACADEMIC_DATA_FILE, "w", encoding="utf-8") as fh:
                json.dump(_DEFAULT_DATA, fh, indent=2)
            return json.loads(json.dumps(_DEFAULT_DATA))
        try:
            with open(ACADEMIC_DATA_FILE, "r", encoding="utf-8") as fh:
                return json.load(fh)
        except Exception:
            return json.loads(json.dumps(_DEFAULT_DATA))


def _save_data(data):
    with _lock:
        with open(ACADEMIC_DATA_FILE, "w", encoding="utf-8") as fh:
            json.dump(data, fh, indent=2, ensure_ascii=False)


def _student_key(program, semester, section):
    return program + "|" + semester + "|" + section


def get_programs():
    return sorted(_load_data().get("programs", {}).keys())


def get_semesters(program):
    prog = _load_data().get("programs", {}).get(program, {})
    return list(prog.get("semesters", {}).keys())


def get_sections(program, semester):
    sems = _load_data().get("programs", {}).get(program, {}).get("semesters", {})
    result = sems.get(semester, {}).get("sections", ["NA"])
    return result if result else ["NA"]


def get_courses(program, semester):
    sems = _load_data().get("programs", {}).get(program, {}).get("semesters", {})
    return sems.get(semester, {}).get("courses", [])


def get_students(program, semester, section):
    key = _student_key(program, semester, section)
    return list(_load_data().get("students", {}).get(key, []))


def get_all_programs_data():
    return _load_data().get("programs", {})


def add_student(program, semester, section, roll_no, enrollment_no, name):
    roll_no = str(roll_no).strip()
    name = str(name).strip()
    if not roll_no or not name:
        return False
    data = _load_data()
    key = _student_key(program, semester, section)
    data.setdefault("students", {})
    students = data["students"].get(key, [])
    if any(s["roll_no"].lower() == roll_no.lower() for s in students):
        return False
    students.append({"roll_no": roll_no, "enrollment_no": str(enrollment_no).strip(), "name": name})
    data["students"][key] = students
    _save_data(data)
    return True


def remove_student(program, semester, section, roll_no):
    data = _load_data()
    key = _student_key(program, semester, section)
    students = data.get("students", {}).get(key, [])
    new_list = [s for s in students if s["roll_no"].lower() != roll_no.lower()]
    if len(new_list) == len(students):
        return False
    data["students"][key] = new_list
    _save_data(data)
    return True


def update_student(program, semester, section, roll_no, **kwargs):
    data = _load_data()
    key = _student_key(program, semester, section)
    students = data.get("students", {}).get(key, [])
    for s in students:
        if s["roll_no"].lower() == roll_no.lower():
            s.update({k: v for k, v in kwargs.items() if k in ("name", "enrollment_no")})
            data["students"][key] = students
            _save_data(data)
            return True
    return False


def add_program(program):
    program = program.strip()
    if not program:
        return False
    data = _load_data()
    if program in data.get("programs", {}):
        return False
    data.setdefault("programs", {})[program] = {"semesters": {}}
    _save_data(data)
    return True


def add_semester(program, semester, sections, courses):
    data = _load_data()
    if program not in data.get("programs", {}):
        return False
    data["programs"][program]["semesters"][semester] = {"sections": list(sections), "courses": list(courses)}
    _save_data(data)
    return True


def import_students_from_excel(program, semester, section, file_path_or_bytes):
    try:
        from openpyxl import load_workbook
    except ImportError:
        return 0, "openpyxl is not installed."
    try:
        if isinstance(file_path_or_bytes, (str, os.PathLike)):
            wb = load_workbook(file_path_or_bytes, read_only=True, data_only=True)
        elif isinstance(file_path_or_bytes, bytes):
            wb = load_workbook(io.BytesIO(file_path_or_bytes), read_only=True, data_only=True)
        else:
            wb = load_workbook(file_path_or_bytes, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            wb.close()
            return 0, "Empty file."
        header_row_idx = None
        headers = []
        for idx, row in enumerate(rows):
            str_row = [str(v).strip().lower() if v else "" for v in row]
            if any("roll" in h for h in str_row):
                header_row_idx = idx
                headers = str_row
                break
        if header_row_idx is None:
            wb.close()
            return 0, "Could not find Roll No header."
        def _fc(*cands):
            for c in cands:
                for i, h in enumerate(headers):
                    if c in h:
                        return i
            return -1
        roll_col = _fc("roll")
        enroll_col = _fc("enrollment", "enroll")
        name_col = _fc("name", "student")
        if roll_col < 0 or name_col < 0:
            wb.close()
            return 0, "Required columns not found."
        imported = skipped = 0
        for row in rows[header_row_idx + 1:]:
            if not row or not row[roll_col]:
                continue
            roll = str(row[roll_col]).strip()
            name = str(row[name_col]).strip() if row[name_col] else ""
            enrl = str(row[enroll_col]).strip() if enroll_col >= 0 and len(row) > enroll_col and row[enroll_col] else ""
            if not roll or not name:
                skipped += 1
                continue
            ok = add_student(program, semester, section, roll, enrl, name)
            if ok:
                imported += 1
            else:
                skipped += 1
        wb.close()
        return imported, "Imported " + str(imported) + " student(s). " + str(skipped) + " skipped."
    except Exception as exc:
        return 0, "Import failed: " + str(exc)


# ── Event / Activity Management ───────────────────────────────────────────────

_DEFAULT_EVENTS = [
    {"name": "Code Champ Competition", "type": "Technical", "date": "2026-09-15"},
    {"name": "Tech Hackathon", "type": "Technical", "date": "2026-09-15"},
    {"name": "Annual Sports Meet", "type": "Sports", "date": "2026-09-15"},
    {"name": "Cultural Fest", "type": "Cultural", "date": "2026-09-15"},
    {"name": "AI & Robotics Workshop", "type": "Workshop", "date": "2026-09-15"},
    {"name": "Inter-College Symposium", "type": "Academic", "date": "2026-09-15"},
]


def get_events() -> list:
    """Return list of approved events/activities."""
    data = _load_data()
    events = data.get("events", [])
    if not events:
        events = list(_DEFAULT_EVENTS)
        data["events"] = events
        _save_data(data)
    return events


def add_event(name: str, event_type: str = "Technical", date_str: str = None) -> bool:
    """Add a new approved event/activity to the master list."""
    name = str(name).strip()
    if not name:
        return False
    from datetime import datetime
    date_str = date_str or datetime.now().strftime("%Y-%m-%d")
    data = _load_data()
    events = data.get("events", [])
    if any(e["name"].lower() == name.lower() for e in events):
        return False
    events.append({"name": name, "type": event_type, "date": date_str})
    data["events"] = events
    _save_data(data)
    return True


def record_event_attendance(event_name: str, roll_no: str,
                            student_name: str = "", remarks: str = "") -> dict:
    """
    Record event/activity attendance for a student automatically.
    Links participation to the student and stores in event_records.
    Returns the recorded event record dict.
    """
    from datetime import datetime
    now_dt = datetime.now()
    event_name = str(event_name).strip()
    roll_no = str(roll_no).strip()

    data = _load_data()
    records = data.get("event_records", [])

    record = {
        "event_name": event_name,
        "roll_no": roll_no,
        "student_name": student_name,
        "date": now_dt.strftime("%Y-%m-%d"),
        "time": now_dt.strftime("%H:%M:%S"),
        "status": "OD",
        "engagement": "Approved Participation",
        "remarks": remarks or "Approved Event Participation"
    }

    # Avoid exact duplicate for same event, roll_no, date
    exists = any(
        r["event_name"].lower() == event_name.lower()
        and r["roll_no"].lower() == roll_no.lower()
        and r.get("date") == record["date"]
        for r in records
    )
    if not exists:
        records.append(record)
        data["event_records"] = records
        _save_data(data)

    return record


def get_event_attendance(event_name: str = None) -> list:
    """Return all recorded event attendance rows (optionally filtered by event_name)."""
    data = _load_data()
    records = data.get("event_records", [])
    if event_name:
        return [r for r in records if r["event_name"].lower() == event_name.lower()]
    return list(records)


def get_student_event_records(roll_no: str) -> list:
    """Return all events attended by student with *roll_no*."""
    data = _load_data()
    records = data.get("event_records", [])
    roll_no = str(roll_no).strip().lower()
    return [r for r in records if str(r.get("roll_no", "")).lower() == roll_no]

